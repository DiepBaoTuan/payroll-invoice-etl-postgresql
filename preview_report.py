from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import REPORTS_DIR


HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
SECTION_FILL = PatternFill("solid", fgColor="E2F0D9")
WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")
TOTAL_FILL = PatternFill("solid", fgColor="D9EAD3")

THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)


def safe_sheet_name(value):
    invalid_chars = ["\\", "/", "*", "[", "]", ":", "?"]
    text = str(value or "Sheet")

    for char in invalid_chars:
        text = text.replace(char, "-")

    text = text.strip()

    if not text:
        text = "Sheet"

    return text[:31]


def get_first_value(df, column_name, default="Missing"):
    if column_name not in df.columns:
        return default

    values = df[column_name].dropna()

    if len(values) == 0:
        return default

    return values.iloc[0]


def build_section_summary(df):
    return (
        df.groupby(
            ["section_type", "week_ending"],
            dropna=False,
            sort=False,
        )
        .agg(
            rows=("employee_number", "count"),
            total_hours=("total_hours", "sum"),
            total_amount=("total_amount", "sum"),
        )
        .reset_index()
    )

def confirmed_bill_difference_rows(df):
    if "adjustment_type" not in df.columns:
        return df.iloc[0:0]

    return df[df["adjustment_type"].eq("bill_difference")]


def bill_difference_candidate_rows(df):
    if "_bill_difference_candidate" not in df.columns:
        return df.iloc[0:0]

    if "adjustment_type" not in df.columns:
        return df.iloc[0:0]

    return df[
        df["_bill_difference_candidate"].eq(True)
        & ~df["adjustment_type"].eq("bill_difference")
    ]


def bill_difference_counts(df):
    confirmed = confirmed_bill_difference_rows(df)
    candidates = bill_difference_candidate_rows(df)

    return len(confirmed), len(candidates)
        
def confirmed_pay_difference_rows(df):
    if "adjustment_ty-pe" not in df.columns:
        return df.iloc[0:0]

    return df[df["adjustment_type"].eq("pay_difference")]

def confirmed_pay_difference_rows(df):
    if "adjustment_type" not in df.columns:
        return df.iloc[0:0]

    return df[df["adjustment_type"].eq("pay_difference")]


def candidate_rows(df):
    if "_pay_difference_candidate" not in df.columns:
        return df.iloc[0:0]

    if "adjustment_type" not in df.columns:
        return df.iloc[0:0]

    return df[
        df["_pay_difference_candidate"].eq(True)
        & ~df["adjustment_type"].eq("pay_difference")
    ]


def pay_difference_counts(df):
    confirmed = confirmed_pay_difference_rows(df)
    candidates = candidate_rows(df)

    return len(confirmed), len(candidates)

def style_header_row(ws, row_number):
    for cell in ws[row_number]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")


def style_range_border(ws, min_row, max_row, min_col, max_col):
    for row in ws.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
    ):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top")


def autosize_columns(ws, max_width=45):
    for column_cells in ws.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = 0

        for cell in column_cells:
            value = cell.value

            if value is None:
                continue

            max_length = max(max_length, len(str(value)))

        ws.column_dimensions[column_letter].width = min(max_length + 2, max_width)


def write_workbook_summary(ws, parsed_invoices):
    ws.title = "Summary"

    headers = [
        "Selection #",
        "Invoice Number",
        "Customer",
        "File Name",
        "Week Endings",
        "Summary Type",
        "Section",
        "Section Week Ending",
        "Rows",
        "Total Hours",
        "Total Amount",
        "Pay Diff Confirmed",
        "Pay Diff Candidates",
        "Bill Diff Confirmed",
        "Bill Diff Candidates",
        "Parsing Seconds",
    ]

    ws.append(headers)
    style_header_row(ws, 1)

    for invoice in parsed_invoices:
        df = invoice["df"]

        invoice_index = invoice["invoice_index"]
        file_name = invoice["file_path"].name
        invoice_number = get_first_value(df, "_invoice_number")
        customer = get_first_value(df, "_customer")

        week_endings = sorted(
            [
                str(value)
                for value in df["week_ending"].dropna().unique()
            ],
            reverse=True,
        )

        total_hours = float(df["total_hours"].fillna(0).sum())
        total_amount = float(df["total_amount"].fillna(0).sum())

        pay_confirmed_count, pay_candidate_count = pay_difference_counts(df)
        bill_confirmed_count, bill_candidate_count = bill_difference_counts(df)

        # Grand total row
        ws.append(
            [
                invoice_index,
                invoice_number,
                customer,
                file_name,
                ", ".join(week_endings),
                "GRAND TOTAL",
                "",
                "",
                len(df),
                round(total_hours, 2),
                round(total_amount, 2),
                pay_confirmed_count,
                pay_candidate_count,
                bill_confirmed_count,
                bill_candidate_count,
                round(float(invoice.get("processing_seconds") or 0), 2),
            ]
        )

        grand_total_row = ws.max_row

        for cell in ws[grand_total_row]:
            cell.font = Font(bold=True)
            cell.fill = TOTAL_FILL

        # Section rows
        section_summary = build_section_summary(df)

        for _, section_row_data in section_summary.iterrows():
            section_type = section_row_data.get("section_type")
            section_week_ending = section_row_data.get("week_ending")

            section_df = df[
                df["section_type"].eq(section_type)
                & df["week_ending"].eq(section_week_ending)
            ]

            section_pay_confirmed_count, section_pay_candidate_count = (
                pay_difference_counts(section_df)
            )
            section_bill_confirmed_count, section_bill_candidate_count = (
                bill_difference_counts(section_df)
            )

            ws.append(
                [
                    invoice_index,
                    invoice_number,
                    customer,
                    file_name,
                    ", ".join(week_endings),
                    "SECTION",
                    section_type,
                    section_week_ending,
                    int(section_row_data.get("rows") or 0),
                    round(float(section_row_data.get("total_hours") or 0), 2),
                    round(float(section_row_data.get("total_amount") or 0), 2),
                    section_pay_confirmed_count,
                    section_pay_candidate_count,
                    section_bill_confirmed_count,
                    section_bill_candidate_count,
                    "",
                ]
            )

            section_row = ws.max_row

            for cell in ws[section_row]:
                cell.fill = SECTION_FILL

        # Blank spacer row between invoices
        ws.append([])

    for row in ws.iter_rows(min_row=2):
        # Total Hours = column 10
        if row[9].value is not None:
            row[9].number_format = "#,##0.00"

        # Total Amount = column 11
        if row[10].value is not None:
            row[10].number_format = "$#,##0.00"

        # Parsing Seconds = column 16
        if row[15].value is not None:
            row[15].number_format = "0.00"

        pay_confirmed_value = row[11].value
        pay_candidate_value = row[12].value
        bill_confirmed_value = row[13].value
        bill_candidate_value = row[14].value

        has_adjustment_rows = (
            isinstance(pay_confirmed_value, int)
            and pay_confirmed_value > 0
        ) or (
            isinstance(pay_candidate_value, int)
            and pay_candidate_value > 0
        ) or (
            isinstance(bill_confirmed_value, int)
            and bill_confirmed_value > 0
        ) or (
            isinstance(bill_candidate_value, int)
            and bill_candidate_value > 0
        )

        if has_adjustment_rows:
            for cell in row:
                if row[5].value == "GRAND TOTAL":
                    cell.fill = WARNING_FILL

    style_range_border(ws, 1, ws.max_row, 1, len(headers))
    autosize_columns(ws)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    
    
def write_section_summary(ws, parsed_invoices):
    ws.title = "By Section"

    headers = [
        "Selection #",
        "Invoice Number",
        "Customer",
        "File Name",
        "Section",
        "Week Ending",
        "Rows",
        "Hours",
        "Total Amount",
    ]

    ws.append(headers)
    style_header_row(ws, 1)

    for invoice in parsed_invoices:
        df = invoice["df"]
        section_summary = build_section_summary(df)

        invoice_index = invoice["invoice_index"]
        invoice_number = get_first_value(df, "_invoice_number")
        customer = get_first_value(df, "_customer")
        file_name = invoice["file_path"].name

        for _, row in section_summary.iterrows():
            ws.append(
                [
                    invoice_index,
                    invoice_number,
                    customer,
                    file_name,
                    row.get("section_type"),
                    row.get("week_ending"),
                    int(row.get("rows") or 0),
                    round(float(row.get("total_hours") or 0), 2),
                    round(float(row.get("total_amount") or 0), 2),
                ]
            )

    for row in ws.iter_rows(min_row=2):
        row[7].number_format = "#,##0.00"
        row[8].number_format = "$#,##0.00"

    style_range_border(ws, 1, ws.max_row, 1, len(headers))
    autosize_columns(ws)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_pay_difference_rows(ws, parsed_invoices):
    ws.title = "Pay Difference Rows"

    headers = [
        "Status",
        "Selection #",
        "Invoice Number",
        "Customer",
        "File Name",
        "Excel Row",
        "Employee Number",
        "First Name",
        "Last Name",
        "Week Ending",
        "Section",
        "Adjustment Type",
        "Pay Rate",
        "Total Hours",
        "Total Amount",
        "Review Notes",
    ]

    ws.append(headers)
    style_header_row(ws, 1)

    for invoice in parsed_invoices:
        df = invoice["df"]

        confirmed = confirmed_pay_difference_rows(df)
        candidates = candidate_rows(df)

        invoice_index = invoice["invoice_index"]
        invoice_number = get_first_value(df, "_invoice_number")
        customer = get_first_value(df, "_customer")
        file_name = invoice["file_path"].name

        for status, rows_df in [
            ("confirmed", confirmed),
            ("candidate", candidates),
        ]:
            if rows_df.empty:
                continue

            for _, row in rows_df.iterrows():
                ws.append(
                    [
                        status,
                        invoice_index,
                        invoice_number,
                        customer,
                        file_name,
                        row.get("_excel_row_number"),
                        row.get("employee_number"),
                        row.get("first_name"),
                        row.get("last_name"),
                        row.get("week_ending"),
                        row.get("section_type"),
                        row.get("adjustment_type"),
                        round(float(row.get("pay_rate") or 0), 2),
                        round(float(row.get("total_hours") or 0), 2),
                        round(float(row.get("total_amount") or 0), 2),
                        "",
                    ]
                )

                current_row = ws.max_row

                for cell in ws[current_row]:
                    if status == "candidate":
                        cell.fill = WARNING_FILL
                    else:
                        cell.fill = SECTION_FILL

    for row in ws.iter_rows(min_row=2):
        row[12].number_format = "$#,##0.00"
        row[13].number_format = "#,##0.00"
        row[14].number_format = "$#,##0.00"

    style_range_border(ws, 1, ws.max_row, 1, len(headers))
    autosize_columns(ws)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_all_rows(ws, parsed_invoices):
    ws.title = "Rows"

    headers = [
        "Selection #",
        "Invoice Number",
        "Customer",
        "File Name",
        "Excel Row",
        "Employee Number",
        "Last Name",
        "First Name",
        "Area",
        "Subarea",
        "Shift",
        "Plant",
        "Position",
        "Section",
        "Rate Type",
        "Adjustment Type",
        "Week Ending",
        "Pay Rate",
        "Weekly Rate",
        "Total Hours",
        "Regular Hours",
        "Overtime Hours",
        "Regular Bill Rate",
        "Overtime Bill Rate",
        "Regular Markup",
        "Overtime Markup",
        "Total Amount",
        "Pay Diff Candidate",
        "Pay Diff Status",
    ]

    ws.append(headers)
    style_header_row(ws, 1)

    for invoice in parsed_invoices:
        df = invoice["df"]

        invoice_index = invoice["invoice_index"]
        invoice_number = get_first_value(df, "_invoice_number")
        customer = get_first_value(df, "_customer")
        file_name = invoice["file_path"].name

        for _, row in df.iterrows():
            if row.get("adjustment_type") == "pay_difference":
                pay_diff_status = "confirmed"
            elif row.get("_pay_difference_candidate") is True:
                pay_diff_status = "candidate"
            else:
                pay_diff_status = ""
                
            ws.append(
                [
                    invoice_index,
                    invoice_number,
                    customer,
                    file_name,
                    row.get("_excel_row_number"),
                    row.get("employee_number"),
                    row.get("last_name"),
                    row.get("first_name"),
                    row.get("area"),
                    row.get("subarea"),
                    row.get("shift"),
                    row.get("plant"),
                    row.get("position"),
                    row.get("section_type"),
                    row.get("rate_type"),
                    row.get("adjustment_type"),
                    row.get("week_ending"),
                    row.get("pay_rate"),
                    row.get("weekly_rate"),
                    row.get("total_hours"),
                    row.get("regular_hours"),
                    row.get("overtime_hours"),
                    row.get("regular_bill_rate"),
                    row.get("overtime_bill_rate"),
                    row.get("regular_markup"),
                    row.get("overtime_markup"),
                    row.get("total_amount"),
                    row.get("_pay_difference_candidate"),
                    pay_diff_status,
                ]
            )

    money_columns = [18, 19, 23, 24, 25, 26, 27]
    hours_columns = [20, 21, 22]

    for row in ws.iter_rows(min_row=2):
        for col_idx in money_columns:
            row[col_idx - 1].number_format = "$#,##0.00"

        for col_idx in hours_columns:
            row[col_idx - 1].number_format = "#,##0.00"

        pay_diff_status = row[28].value

        if pay_diff_status == "candidate":
            for cell in row:
                cell.fill = WARNING_FILL

        elif pay_diff_status == "confirmed":
            for cell in row:
                cell.fill = SECTION_FILL

    style_range_border(ws, 1, ws.max_row, 1, len(headers))
    autosize_columns(ws, max_width=38)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def create_preview_report(parsed_invoices):
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    report_path = REPORTS_DIR / f"invoice_preview_report_{timestamp}.xlsx"

    wb = Workbook()

    summary_ws = wb.active
    write_workbook_summary(summary_ws, parsed_invoices)

    section_ws = wb.create_sheet("By Section")
    write_section_summary(section_ws, parsed_invoices)

    pay_diff_ws = wb.create_sheet("Pay Difference Rows")
    write_pay_difference_rows(pay_diff_ws, parsed_invoices)

    rows_ws = wb.create_sheet("Rows")
    write_all_rows(rows_ws, parsed_invoices)

    wb.save(report_path)

    return report_path